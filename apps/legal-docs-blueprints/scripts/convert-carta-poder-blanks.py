#!/usr/bin/env python3
"""
Convierte los templates de carta_poder con guiones bajos (_______) al formato
con placeholders {key}, usando CARTA-CUBE-ANDRES-HOMBRE.docx como referencia.

Lógica:
- Para cada párrafo del documento target, alinea con el mismo párrafo de la
  referencia. Cada bloque _______ (3+ underscores) se reemplaza con el
  {key} correspondiente que aparece en la referencia en la misma posición.
- Los párrafos de la referencia que no tienen {key} (ej. líneas F.____ de
  firma manual) se ignoran y los underscores del target quedan intactos.
- {edadAndres} se sustituye según el archivo:
    CUBE-ANDRES-*       → edadAndres
    CUBE-DON ALEX-*     → edadAlex
    RDBE-DON ALEX-*     → edadAlex
    RDBE-RICHARD-*      → edadRichard

Sobreescribe los archivos in-place (sin backup).
"""
from pathlib import Path
import re
import sys

from docx import Document
from docx.document import Document as _Doc
from docx.oxml.ns import qn
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph

DIR = Path(__file__).parent.parent / "templates" / "carta_poder"
REFERENCE = "CARTA-CUBE-ANDRES-HOMBRE.docx"
TARGETS = [
    "CARTA-CUBE-ANDRES-MUJER.docx",
    "CARTA-CUBE-DON ALEX-HOMBRE.docx",
    "CARTA-CUBE-DON ALEX-MUJER.docx",
    "CARTA-RDBE-DON ALEX-HOMBRE.docx",
    "CARTA-RDBE-DON ALEX-MUJER.docx",
    "CARTA-RDBE-RICHARD-HOMBRE.docx",
    "CARTA-RDBE-RICHARD-MUJER.docx",
]

EDAD_MAP = {
    "CUBE-ANDRES":   "edadAndres",
    "CUBE-DON ALEX": "edadAlex",
    "RDBE-DON ALEX": "edadAlex",
    "RDBE-RICHARD":  "edadRichard",
}

KEY_RE = re.compile(r"\{[a-zA-Z][a-zA-Z0-9_]*\}")
UND_RE = re.compile(r"_{3,}")


def walk_paragraphs(parent):
    if isinstance(parent, _Doc):
        elem = parent.element.body
    elif isinstance(parent, _Cell):
        elem = parent._tc
    else:
        return
    for child in elem.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, parent)
        elif child.tag == qn("w:tbl"):
            for row in Table(child, parent).rows:
                for cell in row.cells:
                    yield from walk_paragraphs(cell)


def replace_blocks(para, keys):
    """
    Reemplaza bloques _____ del párrafo con la lista `keys` en orden.
    Maneja bloques que se extienden por múltiples runs.
    Retorna cantidad de reemplazos, o -1 si la cantidad de bloques no
    coincide con la cantidad de keys.
    """
    runs = list(para.runs)
    if not runs:
        return 0

    full_text = "".join(r.text for r in runs)
    blocks = list(UND_RE.finditer(full_text))

    if not blocks:
        return 0
    if len(blocks) != len(keys):
        return -1

    # Offsets originales por run
    run_starts = []
    pos = 0
    for r in runs:
        run_starts.append(pos)
        pos += len(r.text)

    new_text = [r.text for r in runs]

    # Procesar bloques de derecha a izquierda para que los offsets
    # de la izquierda permanezcan válidos tras cada edición.
    for block, key in reversed(list(zip(blocks, keys))):
        bs, be = block.start(), block.end()
        affected = [
            ri for ri, rs in enumerate(run_starts)
            if rs < be and rs + len(runs[ri].text) > bs
        ]
        for idx, ri in enumerate(affected):
            rs = run_starts[ri]
            re_ = rs + len(runs[ri].text)
            local_s = max(rs, bs) - rs
            local_e = min(re_, be) - rs
            cur = new_text[ri]
            replacement = key if idx == 0 else ""
            new_text[ri] = cur[:local_s] + replacement + cur[local_e:]

    for r, t in zip(runs, new_text):
        if r.text != t:
            r.text = t

    return len(blocks)


def edad_role(filename):
    for pat, key in EDAD_MAP.items():
        if pat in filename:
            return key
    return "edadAndres"


def main():
    ref_path = DIR / REFERENCE
    if not ref_path.exists():
        print(f"❌ No existe: {ref_path}", file=sys.stderr)
        sys.exit(1)

    ref_doc = Document(str(ref_path))
    ref_paras = list(walk_paragraphs(ref_doc))
    print(f"📘 Referencia: {REFERENCE} ({len(ref_paras)} párrafos)\n")

    overall_warnings = 0
    for tname in TARGETS:
        tpath = DIR / tname
        if not tpath.exists():
            print(f"⚠  No existe: {tname}\n")
            continue

        tdoc = Document(str(tpath))
        tparas = list(walk_paragraphs(tdoc))
        edad_key = edad_role(tname)

        print(f"📄 {tname}  (edad → {edad_key})")
        if len(tparas) != len(ref_paras):
            print(f"   ℹ  párrafos: ref={len(ref_paras)} tgt={len(tparas)}")

        total_replaced = 0
        warnings = []
        for i, (rp, tp) in enumerate(zip(ref_paras, tparas)):
            ref_keys_raw = KEY_RE.findall(rp.text)
            if not ref_keys_raw:
                continue  # párrafo sin placeholders → no tocar (ej. F.____ firma)

            # Sustituir {edadAndres} por la key correspondiente al archivo
            ref_keys = [
                "{" + edad_key + "}" if k == "{edadAndres}" else k
                for k in ref_keys_raw
            ]

            result = replace_blocks(tp, ref_keys)
            if result == -1:
                blanks_in_tgt = len(UND_RE.findall(tp.text))
                warnings.append(
                    f"   ⚠  párrafo #{i}: ref={len(ref_keys)} keys vs "
                    f"tgt={blanks_in_tgt} blanks → SE OMITIÓ ese párrafo"
                )
            else:
                total_replaced += result

        tdoc.save(str(tpath))
        print(f"   ✅ reemplazos aplicados: {total_replaced}")
        for w in warnings:
            print(w)
            overall_warnings += 1
        print()

    if overall_warnings:
        print(f"⚠  {overall_warnings} párrafo(s) con desbalance — revisalos manualmente.")
        sys.exit(2)
    print("✅ Sin warnings — alineamiento correcto en todos los archivos procesados.")


if __name__ == "__main__":
    main()
